from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
import json

path = Path('test-results-visits/workbook-validation.xlsx')
with ZipFile(path) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith(('.xml', '.rels')):
            ET.fromstring(archive.read(name))
book = load_workbook(path)
assert book.sheetnames == ['Resumo mensal', 'Visitas', 'Diário', 'Semanal', 'PDVs', 'Dashboard', 'Metodologia']
assert len(book['Dashboard']._charts) == 2
assert book['Visitas']['P2'].value.startswith('=HYPERLINK') and book['Visitas']['P2'].data_type == 's'
assert book['Visitas']['Q2'].value is None
assert book['Resumo mensal']['B5'].value == 3
assert book['Resumo mensal']['B6'].value == 2
assert book['Diário']['B2'].value == 2
assert book['Diário']['C2'].value == 1
assert book['Diário']['K3'].value == 2
assert book['Dashboard']._charts[1].ser[0].val.numRef.f == "'Diário'!$K$2:$K$32"
model = json.loads(Path('test-results-visits/workbook-model.json').read_text())
for row, daily in enumerate(model['daily'], start=2):
    assert book['Diário'].cell(row, 10).value == daily['projectedPdvs']
    assert book['Diário'].cell(row, 11).value == daily['cumulativeActivePdvs']
print('Validated real XLSX: seven sheets, two editable charts, report values, UTF-8 text and safe literal cells.')
