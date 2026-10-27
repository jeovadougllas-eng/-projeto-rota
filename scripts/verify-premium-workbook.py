from openpyxl import load_workbook
from zipfile import ZipFile
from xml.etree import ElementTree as ET
for tier in ['pro','plus']:
    path=f'test-results-premium/{tier}-workbook.xlsx'
    with ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith(('.xml','.rels')):ET.fromstring(z.read(name))
    b=load_workbook(path)
    assert len(b['Dashboard']._charts)==2
    if tier=='pro':
        assert 'Comparação' in b.sheetnames
        assert b['Comparação']['C2'].value==1
        assert b['Comparação']['C3'].value==2
        assert b['Comparação']['D3'].value==1
    else:
        assert 'Comparação' not in b.sheetnames
        assert all(b['Diário'].cell(i,9).value is None and b['Diário'].cell(i,10).value is None for i in range(2,33))
print('Verified Plus and Pro XLSX files, real editable charts, tier contents and distinct PDV comparison.')
